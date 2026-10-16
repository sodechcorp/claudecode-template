# -*- coding: utf-8 -*-
"""Excel _meta シートへのメタデータ読み書き"""

import hashlib
import json
from pathlib import Path as _Path

from openpyxl import load_workbook

META_SHEET = "_meta"

# _meta に保存しない大きなキー（describe は生の Salesforce レスポンスで巨大、field_usage は毎回再取得）
_EXCLUDE_KEYS = {"describe", "field_usage"}


# Excel セルの文字数上限は 32,767 文字なので分割して保存する
_CHUNK = 30000


def read_meta(source_file: str) -> dict | None:
    """既存 Excel の _meta シートからメタデータを読み込む。
    _meta シートが無いときは None を返す。読めないときは例外を上げる（None は新規作成として扱われ、改版履歴が消える）。"""
    try:
        wb = load_workbook(source_file, read_only=True, data_only=True)
        if META_SHEET not in wb.sheetnames:
            print(f"  [INFO] _meta シートなし: {source_file}")
            wb.close()
            return None
        ws = wb[META_SHEET]
        chunks = []
        for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
            if row[0] is not None:
                chunks.append(str(row[0]))
        wb.close()
        return json.loads("".join(chunks))
    except Exception as e:
        raise RuntimeError(f"_meta の読込に失敗: {source_file}: {e}（新規作成にすると改版履歴が消えるため書かずに止めた。"
                           f"ファイルは動かさずに報告する。作り直すなら担当者が別の場所へ移して再実行する）") from e


def write_meta(wb, data: dict):
    """Workbook に _meta シートを追加（非表示）。既存シートは上書き。
    JSON が 30,000 字を超える場合は複数行に分割して保存する。"""
    if META_SHEET in wb.sheetnames:
        del wb[META_SHEET]
    ws = wb.create_sheet(META_SHEET)
    ws.sheet_state = "hidden"
    json_str = json.dumps(data, ensure_ascii=False, default=str)
    for i, start in enumerate(range(0, len(json_str), _CHUNK)):
        ws.cell(row=i + 1, column=1, value=json_str[start:start + _CHUNK])


def strip_meta(obj: dict) -> dict:
    """保存不要な大きなキーを除去した軽量メタデータを返す"""
    return {k: v for k, v in obj.items() if k not in _EXCLUDE_KEYS}


def compute_source_hash(paths: list) -> str:
    """複数ファイル／ディレクトリの SHA256 ハッシュを計算する。
    ディレクトリが渡された場合は配下の全ファイルを再帰展開する。
    パスでソートして順序を安定させる。
    """
    h = hashlib.sha256()
    all_files: list[str] = []
    for raw in paths:
        p = _Path(str(raw))
        if p.is_dir():
            all_files.extend(str(f) for f in p.rglob("*") if f.is_file())
        elif p.exists():
            all_files.append(str(p))
    for f in sorted(all_files):
        h.update(_Path(f).read_bytes())
    return h.hexdigest()


def get_stored_hash(source_file: str) -> str | None:
    """既存 Excel の _meta から source_hash を取得する。なければ None。"""
    meta = read_meta(source_file)
    return meta.get("source_hash") if meta else None
